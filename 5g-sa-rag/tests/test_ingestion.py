"""
tests/test_ingestion.py

Unit tests for the ingestion pipeline.
Run: pytest tests/ -v
"""
import pytest
from pathlib import Path
from docx import Document
import tempfile
import os

from ingestion.entity_extractor import EntityExtractor
from ingestion.chunker import SectionAwareChunker
from ingestion.parser import DocxParser, detect_carrier, detect_doc_type


# ── Entity Extractor tests ────────────────────────────────────────────────────

class TestEntityExtractor:
    def setup_method(self):
        self.ext = EntityExtractor()

    def test_extracts_req_ids(self):
        text = "The system shall implement REQ-SA-001 and REQ-QOS-003."
        entities = self.ext.extract(text)
        assert "REQ-SA-001" in entities.req_ids
        assert "REQ-QOS-003" in entities.req_ids

    def test_extracts_tc_ids(self):
        text = "Test case TC-SA-042 verifies the handover procedure."
        entities = self.ext.extract(text)
        assert "TC-SA-042" in entities.tc_ids

    def test_extracts_features(self):
        text = "The handover procedure shall complete within 50ms for URLLC services."
        entities = self.ext.extract(text)
        assert "Handover" in entities.features
        assert "URLLC" in entities.features

    def test_extracts_nr_entities(self):
        text = "The AMF shall authenticate the UE via AUSF using 5G-AKA."
        entities = self.ext.extract(text)
        assert "AMF" in entities.nr_entities
        assert "AUSF" in entities.nr_entities

    def test_extracts_req_to_tc_links(self):
        text = "REQ-SA-001 is verified by TC-SA-001 as part of the test suite."
        links = self.ext.extract_req_to_tc_links(text)
        assert ("REQ-SA-001", "TC-SA-001") in links

    def test_no_false_positives(self):
        text = "The system shall provide adequate performance."
        entities = self.ext.extract(text)
        assert len(entities.req_ids) == 0
        assert len(entities.tc_ids) == 0


# ── Carrier detection tests ───────────────────────────────────────────────────

class TestCarrierDetection:
    def test_detects_verizon(self):
        assert detect_carrier("This document is for Verizon 5G SA deployment") == "Verizon"

    def test_detects_tmobile(self):
        assert detect_carrier("T-Mobile network requirements") == "T-Mobile"

    def test_detects_att(self):
        assert detect_carrier("AT&T carrier requirements") == "AT&T"

    def test_unknown(self):
        assert detect_carrier("Generic 5G requirements document") == "Unknown"


# ── Doc type detection ────────────────────────────────────────────────────────

class TestDocTypeDetection:
    def test_detects_test_plan_by_name(self):
        assert detect_doc_type("", "Verizon_test_plan_v2") == "test_plan"

    def test_detects_requirements_by_name(self):
        assert detect_doc_type("", "ATT_requirements_5GSA") == "requirements"

    def test_detects_requirements_by_content(self):
        text = "REQ-SA-001 The AMF shall... REQ-SA-002 The SMF shall..."
        assert detect_doc_type(text, "unknown_doc") == "requirements"


# ── Chunker tests ─────────────────────────────────────────────────────────────

class TestChunker:
    def setup_method(self):
        self.chunker = SectionAwareChunker(max_tokens=100, overlap_tokens=10)

    def test_short_section_single_chunk(self):
        from ingestion.parser import DocSection
        section = DocSection(
            doc_name="test_doc", doc_type="requirements", carrier="Verizon",
            heading="Overview", heading_level=1, content="Short content here.",
            section_path=["Overview"],
        )
        chunks = self.chunker.chunk_section(section, 0)
        assert len(chunks) == 1
        assert chunks[0].text == "Short content here."

    def test_chunk_preserves_metadata(self):
        from ingestion.parser import DocSection
        section = DocSection(
            doc_name="my_doc", doc_type="test_plan", carrier="T-Mobile",
            heading="Auth Tests", heading_level=2, content="Test content.",
            section_path=["Test Plan", "Auth Tests"],
        )
        chunks = self.chunker.chunk_section(section, 5)
        assert chunks[0].carrier == "T-Mobile"
        assert chunks[0].doc_name == "my_doc"
        assert chunks[0].heading == "Auth Tests"

    def test_table_becomes_separate_chunk(self):
        from ingestion.parser import DocSection
        section = DocSection(
            doc_name="doc", doc_type="requirements", carrier="AT&T",
            heading="QoS Table", heading_level=1,
            content="Some text.",
            tables=[[["Req ID", "Priority"], ["REQ-001", "High"]]],
            section_path=["QoS"],
        )
        chunks = self.chunker.chunk_section(section, 0)
        types = [c.chunk_type for c in chunks]
        assert "table" in types
        assert "text" in types

    def test_embed_text_includes_context(self):
        from ingestion.parser import DocSection
        section = DocSection(
            doc_name="vz_req", doc_type="requirements", carrier="Verizon",
            heading="Handover", heading_level=1, content="HO content.",
            section_path=["Handover"],
        )
        chunks = self.chunker.chunk_section(section, 0)
        embed_text = chunks[0].to_embed_text()
        assert "Verizon" in embed_text
        assert "requirements" in embed_text
        assert "vz_req" in embed_text


# ── Parser tests (requires creating a temp .docx) ────────────────────────────

class TestDocxParser:
    def _make_docx(self, tmp_dir: str) -> str:
        """Create a minimal test .docx file."""
        doc = Document()
        doc.add_heading("Verizon 5G SA Requirements", level=1)
        doc.add_paragraph("This document covers Verizon carrier requirements.")
        doc.add_heading("QoS Requirements", level=2)
        doc.add_paragraph(
            "REQ-QOS-001: The network shall support 5QI mapping. "
            "This is verified by TC-QOS-001."
        )
        path = os.path.join(tmp_dir, "Verizon_requirements.docx")
        doc.save(path)
        return path

    def test_parse_returns_sections(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = self._make_docx(tmp)
            parser = DocxParser()
            result = parser.parse(path)
            assert result.carrier == "Verizon"
            assert result.doc_type == "requirements"
            assert len(result.sections) >= 1

    def test_parse_detects_carrier(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = self._make_docx(tmp)
            result = DocxParser().parse(path)
            assert result.carrier == "Verizon"
