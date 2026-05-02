"""
ingestion/parser.py

Parses .docx files into structured sections, preserving heading hierarchy.
Handles both requirements docs and test plan docs.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterator

from docx import Document
from docx.oxml.ns import qn


# ─────────────────────────────────────────────────────────────────────────────
#  Data models
# ─────────────────────────────────────────────────────────────────────────────

@dataclass
class DocSection:
    """A logical section extracted from a .docx file."""
    doc_name: str           # source filename (without extension)
    doc_type: str           # "requirements" | "test_plan" | "unknown"
    carrier: str            # e.g. "Verizon", "T-Mobile", "AT&T" — inferred
    heading: str            # section heading text
    heading_level: int      # 1 = top-level, 2 = sub, etc.
    content: str            # full text of the section
    tables: list[list[list[str]]] = field(default_factory=list)  # extracted tables
    page_hint: int = 0      # approximate page number (best-effort)
    section_path: list[str] = field(default_factory=list)  # breadcrumb


@dataclass
class ParsedDocument:
    """Top-level result of parsing a single .docx file."""
    file_path: str
    doc_name: str
    doc_type: str
    carrier: str
    sections: list[DocSection]
    metadata: dict


# ─────────────────────────────────────────────────────────────────────────────
#  Carrier detection
# ─────────────────────────────────────────────────────────────────────────────

CARRIER_PATTERNS = {
    "Verizon":  r"\b(verizon|vzw)\b",
    "T-Mobile": r"\b(t-?mobile|tmobile|tmo)\b",
    "AT&T":     r"\b(at&t|att)\b",
    "Dish":     r"\b(dish|boost)\b",
    "US-Cellular": r"\b(us.?cellular|uscc)\b",
}

def detect_carrier(text: str) -> str:
    """Infer carrier from document text. Returns 'Unknown' if none found."""
    lower = text.lower()
    for carrier, pattern in CARRIER_PATTERNS.items():
        if re.search(pattern, lower, re.IGNORECASE):
            return carrier
    return "Unknown"


# ─────────────────────────────────────────────────────────────────────────────
#  Doc type detection
# ─────────────────────────────────────────────────────────────────────────────

def detect_doc_type(text: str, filename: str) -> str:
    name_lower = filename.lower()
    if any(k in name_lower for k in ["test_plan", "testplan", "test-plan", "tp_"]):
        return "test_plan"
    if any(k in name_lower for k in ["req", "requirement", "spec"]):
        return "requirements"

    # Fall back to content signals
    req_signals = len(re.findall(r"\bREQ[-_]\w+", text))
    tc_signals  = len(re.findall(r"\bTC[-_]\w+|\bTest Case\b", text, re.IGNORECASE))
    if tc_signals > req_signals:
        return "test_plan"
    if req_signals > 0:
        return "requirements"
    return "unknown"


# ─────────────────────────────────────────────────────────────────────────────
#  Table extractor
# ─────────────────────────────────────────────────────────────────────────────

def extract_table(table) -> list[list[str]]:
    """Convert a docx Table object to a 2D list of strings."""
    rows = []
    for row in table.rows:
        cells = [cell.text.strip() for cell in row.cells]
        rows.append(cells)
    return rows


# ─────────────────────────────────────────────────────────────────────────────
#  Core parser
# ─────────────────────────────────────────────────────────────────────────────

class DocxParser:
    """
    Parses a .docx file into structured DocSection objects.

    Strategy:
      - Walk paragraphs in document order
      - When a Heading style is encountered, start a new section
      - Accumulate paragraphs + tables until the next heading
      - Maintain a breadcrumb stack for nested headings
    """

    HEADING_STYLE_RE = re.compile(r"^Heading\s+(\d+)$", re.IGNORECASE)

    def parse(self, file_path: str | Path) -> ParsedDocument:
        path = Path(file_path)
        doc = Document(str(path))

        full_text = "\n".join(p.text for p in doc.paragraphs)
        carrier  = detect_carrier(full_text)
        doc_type = detect_doc_type(full_text, path.stem)

        sections: list[DocSection] = []
        breadcrumb: list[str] = []      # stack of heading texts by level

        current_heading       = "Preamble"
        current_heading_level = 0
        current_content_parts: list[str] = []
        current_tables: list[list[list[str]]] = []

        def flush_section():
            nonlocal current_heading, current_heading_level, current_content_parts, current_tables
            content = "\n".join(current_content_parts).strip()
            if content or current_tables:
                sections.append(DocSection(
                    doc_name=path.stem,
                    doc_type=doc_type,
                    carrier=carrier,
                    heading=current_heading,
                    heading_level=current_heading_level,
                    content=content,
                    tables=list(current_tables),
                    section_path=list(breadcrumb),
                ))
            current_content_parts = []
            current_tables = []

        # Iterate body elements in order (paragraphs + tables interleaved)
        body = doc.element.body
        for child in body:
            tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag

            if tag == "p":
                para_style = child.find(f".//{qn('w:pStyle')}")
                style_name = para_style.get(qn("w:val"), "") if para_style is not None else ""
                text = "".join(
                    node.text or ""
                    for node in child.iter()
                    if node.tag == qn("w:t")
                ).strip()

                m = self.HEADING_STYLE_RE.match(style_name)
                if m:
                    flush_section()
                    level = int(m.group(1))
                    # Update breadcrumb
                    if level <= len(breadcrumb):
                        breadcrumb = breadcrumb[: level - 1]
                    breadcrumb.append(text)
                    current_heading       = text
                    current_heading_level = level
                else:
                    if text:
                        current_content_parts.append(text)

            elif tag == "tbl":
                # Find matching Table object by index
                tbl_elements = [c for c in body if c.tag.split("}")[-1] == "tbl"]
                idx = tbl_elements.index(child) if child in tbl_elements else -1
                if idx >= 0 and idx < len(doc.tables):
                    current_tables.append(extract_table(doc.tables[idx]))

        flush_section()  # Don't forget last section

        return ParsedDocument(
            file_path=str(path),
            doc_name=path.stem,
            doc_type=doc_type,
            carrier=carrier,
            sections=sections,
            metadata={
                "total_sections": len(sections),
                "has_tables": any(s.tables for s in sections),
                "carrier": carrier,
                "doc_type": doc_type,
            },
        )

    def parse_folder(self, folder: str | Path) -> Iterator[ParsedDocument]:
        """Parse all .docx files in a folder."""
        for p in Path(folder).glob("**/*.docx"):
            yield self.parse(p)
